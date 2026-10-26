#!/usr/bin/env python3
"""Build the scenario dataset workbook from the engine export.

    node tools/test/scenario_export.js dataset.json
    python3 tools/build_scenario_xlsx.py dataset.json out.xlsx

Every number in the Outcomes, Sequence and Parameters sheets is engine output at the
scenario's own issue time. Nothing is transcribed from a scenario's description.
"""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

src, dst = sys.argv[1], sys.argv[2]
D = json.load(open(src))

# The prototype's tokens, carried into the workbook.
PANEL, ACC, INK, MUT, LINE = "0D0F12", "79A7E3", "14181D", "55616F", "D6DEE8"
OK, WATCH, HIGH, CRIT = "3FB27F", "E2B93B", "EE8A2A", "E5484D"
UI, MONO = "Barlow", "IBM Plex Mono"

head_font = Font(name=UI, bold=True, size=9, color=ACC)
head_fill = PatternFill("solid", fgColor=PANEL)
body_font = Font(name=UI, size=9.5, color=INK)
mono_font = Font(name=MONO, size=9.5, color=INK)
thin = Side(style="thin", color=LINE)
edge = Border(left=thin, right=thin, top=thin, bottom=thin)

def sheet(wb, name, rows, cols=None, widths=None, mono_cols=(), freeze="A2", note=None):
    ws = wb.create_sheet(name)
    if not rows:
        ws["A1"] = "no rows"; return ws
    cols = cols or list(rows[0].keys())
    r0 = 1
    if note:
        ws.cell(1, 1, note).font = Font(name=UI, size=9, italic=True, color=MUT)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
        r0 = 3
    for j, c in enumerate(cols, 1):
        cell = ws.cell(r0, j, c.replace("_", " "))
        cell.font = head_font; cell.fill = head_fill
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell.border = edge
    for i, row in enumerate(rows, r0 + 1):
        for j, c in enumerate(cols, 1):
            v = row.get(c, "")
            cell = ws.cell(i, j, v)
            cell.font = mono_font if (c in mono_cols or isinstance(v, (int, float))) else body_font
            cell.alignment = Alignment(vertical="top", wrap_text=(isinstance(v, str) and len(str(v)) > 60))
            cell.border = edge
            if i % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F4F7FB")
    for j, c in enumerate(cols, 1):
        w = (widths or {}).get(c)
        if not w:
            w = min(46, max(10, len(c) + 2, *(len(str(r.get(c, ""))) for r in rows[:80])))
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[r0].height = 30
    ws.freeze_panes = ws[freeze if r0 == 1 else f"A{r0+1}"]
    ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(cols))}{r0+len(rows)}"
    return ws

wb = Workbook(); wb.remove(wb.active)

# ── 1. Readme ───────────────────────────────────────────────────────────────
ws = wb.create_sheet("Readme")
lines = [
    ("Cascade Intelligence — scenario dataset", Font(name=UI, bold=True, size=16, color=INK)),
    ("Twenty natural-calamity scenarios added to the simulator, with what the engine does with each one.", Font(name=UI, size=10.5, color=MUT)),
    ("", None),
    ("How this was produced", Font(name=UI, bold=True, size=11, color=INK)),
    ("Every row in Outcomes, Sequence and Parameters is engine output, produced by running each", Font(name=UI, size=10, color=INK)),
    ("city x scenario through the 200-member ensemble and reading the result at that scenario's own", Font(name=UI, size=10, color=INK)),
    ("issue time. Nothing is transcribed from a scenario's written description.", Font(name=UI, size=10, color=INK)),
    ("    node tools/test/scenario_export.js dataset.json", Font(name=MONO, size=9.5, color=MUT)),
    ("    python3 tools/build_scenario_xlsx.py dataset.json out.xlsx", Font(name=MONO, size=9.5, color=MUT)),
    ("", None),
    ("What is real and what is simulated", Font(name=UI, bold=True, size=11, color=INK)),
    ("Real: the facilities, streets, rivers and building footprints of both cities (OpenStreetMap).", Font(name=UI, size=10, color=INK)),
    ("Simulated: every sensor reading, citizen report and outcome. These are modeled scenarios, not", Font(name=UI, size=10, color=INK)),
    ("records of specific events; the one reconstruction in the set (Nagpur, 23 Sep 2023) takes only", Font(name=UI, size=10, color=INK)),
    ("its rainfall timing and totals from reported figures.", Font(name=UI, size=10, color=INK)),
    ("", None),
    ("Sheets", Font(name=UI, bold=True, size=11, color=INK)),
    ("Scenarios    the catalogue as declared: hazard, city, provenance, default issue time.", Font(name=UI, size=10, color=INK)),
    ("Outcomes     what the engine computes for each city x scenario. One row per run.", Font(name=UI, size=10, color=INK)),
    ("Sequence     the ordered chain of CRITICAL crossings per run: facility, system, hour.", Font(name=UI, size=10, color=INK)),
    ("Parameters   the hidden 'truth' parameter set each forecast is scored against.", Font(name=UI, size=10, color=INK)),
    ("", None),
    ("Reading Outcomes", Font(name=UI, bold=True, size=11, color=INK)),
    ("cascade_R         stress arriving through learned dependencies / stress arriving from the hazard.", Font(name=UI, size=10, color=INK)),
    ("                  At or above 1 the cascade sustains itself and stopping the hazard no longer stops it.", Font(name=UI, size=10, color=INK)),
    ("system_order      the order systems first reach CRITICAL. This is the cascade, in one cell.", Font(name=UI, size=10, color=INK)),
    ("act_by_h          hours after the issue time by which the crews must start to keep 80% of the", Font(name=UI, size=10, color=INK)),
    ("                  recommended plan's benefit. Blank means no plan in the catalogue helps enough.", Font(name=UI, size=10, color=INK)),
    ("", None),
    ("Known limitation", Font(name=UI, bold=True, size=11, color=CRIT)),
    ("The Raleigh city model contains no telecom facilities, so peak_telecom is 0 in every Raleigh run", Font(name=UI, size=10, color=INK)),
    ("and no Raleigh scenario can show a telecom failure. This is a gap in that city's harvested data,", Font(name=UI, size=10, color=INK)),
    ("not in the scenarios. Nagpur has five telecom sites and they respond normally.", Font(name=UI, size=10, color=INK)),
]
for i, (t, f) in enumerate(lines, 1):
    c = ws.cell(i, 1, t)
    if f: c.font = f
ws.column_dimensions["A"].width = 112
ws.sheet_view.showGridLines = False

# ── 2–5. Data sheets ────────────────────────────────────────────────────────
sheet(wb, "Scenarios", D["scenarios"],
      widths={"blurb": 95, "label": 36, "id": 15, "set": 22, "provenance": 24},
      mono_cols=("id",),
      note="The catalogue as declared in src/js/engine/scenarios.js and scenarios_natural.js.")

out_cols = list(D["outcomes"][0].keys())
sheet(wb, "Outcomes", D["outcomes"],
      cols=out_cols, widths={"label": 36, "id": 15, "system_order": 34, "recommended_plan": 34, "set": 10, "city": 9},
      mono_cols=("id",),
      note="One row per city x scenario. Every figure is computed by the engine at the scenario's issue time.")

sheet(wb, "Sequence", D["sequence"],
      widths={"facility": 38, "id": 15, "city": 9},
      mono_cols=("id",),
      note="The ordered chain of CRITICAL crossings. 'observed' crossed before the issue time; 'forecast' is ahead of it.")

sheet(wb, "Parameters", D["params"],
      widths={"label": 36, "id": 15, "parameter": 16},
      mono_cols=("id", "parameter"),
      note="The hidden truth: the single parameter set the ensemble is scored against for each scenario.")

# Colour cascade_R and the reduction column on the Outcomes sheet.
ws = wb["Outcomes"]
hdr = {c.value: c.column for c in ws[3]}
for row in ws.iter_rows(min_row=4, max_row=3 + len(D["outcomes"])):
    r = row[hdr["cascade R"] - 1]
    if isinstance(r.value, (int, float)):
        r.font = Font(name=MONO, size=9.5, bold=r.value >= 1,
                      color=CRIT if r.value >= 1 else HIGH if r.value >= 0.6 else INK)
    p = row[hdr["plan reduction pct"] - 1]
    if isinstance(p.value, (int, float)):
        p.font = Font(name=MONO, size=9.5, color=OK if p.value >= 20 else INK)

wb.save(dst)
print("wrote", dst)
